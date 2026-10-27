"""
Автоматизация жизненного цикла реквизита и технического оборудования
(на примере Астраханской государственной филармонии).

Функции: учет, выдача/возврат, бронирование, инвентаризация, контроль состояния,
планирование закупок, отчеты (Excel/PDF), уведомления и фотофиксация возврата реквизита.

Запуск: pip install -r requirements.txt && python app.py  ->  http://127.0.0.1:5000
Первый вход: admin / admin (смените пароль в разделе «Профиль»).
Переменные окружения: см. README.md
"""
import hmac
import io
import math
import os
import re
import secrets
import smtplib
import sqlite3
import threading
import time
import json
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from functools import wraps
from xml.sax.saxutils import escape as xml_escape

from flask import (Flask, g, request, redirect, url_for, session, render_template, flash,
                   abort, send_file, send_from_directory, has_request_context)
from jinja2 import DictLoader
from markupsafe import Markup
from wear import wear, wear_description, wear_level
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB = os.environ.get("DB_PATH", os.path.join(BASE_DIR, "philharmonia.db"))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads", "returns")
ALLOWED_PHOTO_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}
REMIND_DAYS = int(os.environ.get("REMIND_DAYS", "3"))       # за сколько дней предупреждать о возврате
DEFAULT_ISSUE_DAYS = int(os.environ.get("DEFAULT_ISSUE_DAYS", "7"))
CHECK_PERIOD_DAYS = 180  # межремонтный/проверочный интервал
INV_PREFIX = "АГФ-"      # префикс инвентарных номеров

OK_STATE = "Исправно"
SYSTEM_STATES = ["Исправно", "Требует ремонта", "В ремонте", "Списано"]  # на них опирается логика
DEFAULT_CATEGORIES = ["Музыкальные инструменты", "Звук", "Свет", "Сценическое оборудование",
                      "Реквизит", "Костюмы", "Прочее"]
DEFAULT_LOCATIONS = ["Большой зал", "Склад звука", "Склад сцены", "Склад света",
                     "Аппаратная", "Склад реквизита"]
AVAIL_STATUSES = ["В наличии", "Выдан частично", "Выдан", "В ремонте", "Списан", "Отсутствует"]
NO_EVENT = "Хозяйственные нужды (без мероприятия)"

ROLES = {"admin": "Администратор", "responsible": "Ответственный", "user": "Пользователь"}
STAFF = {"admin", "responsible"}
PERMS = {  # право -> роли
    "view": set(ROLES),
    "edit_items": STAFF,      # карточки, акты проверки
    "operations": STAFF,      # выдача/возврат, брони, инвентаризация, закупки
    "dictionaries": STAFF,
    "reports": STAFF,
    "admin": {"admin"},       # управление пользователями
}


def _secret_key():
    """Ключ подписи сессий: из окружения, иначе генерируется один раз и хранится в .secret_key."""
    env = os.environ.get("SECRET_KEY")
    if env:
        return env
    path = os.path.join(BASE_DIR, ".secret_key")
    try:
        with open(path, encoding="ascii") as fh:
            val = fh.read().strip()
            if val:
                return val
    except OSError:
        pass
    val = secrets.token_hex(32)
    try:
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "w", encoding="ascii") as fh:
            fh.write(val)
    except OSError:
        pass
    return val


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # лимит запроса с фотографиями возврата
app.secret_key = _secret_key()
app.config.update(SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_HTTPONLY=True,
                  PERMANENT_SESSION_LIFETIME=timedelta(hours=8))
ADMIN_LOGIN = os.environ.get("ADMIN_LOGIN", "admin")
ADMIN_PASS = os.environ.get("ADMIN_PASS", "admin")

# ---------------------------------------------------------------- БД
SCHEMA = """
CREATE TABLE IF NOT EXISTS items(
  id INTEGER PRIMARY KEY AUTOINCREMENT, inv_no TEXT UNIQUE, name TEXT NOT NULL,
  category TEXT, location TEXT, qty INTEGER DEFAULT 1, min_qty INTEGER DEFAULT 0,
  status TEXT DEFAULT 'Исправно', price REAL DEFAULT 0, bought TEXT,
  life_years REAL DEFAULT 5, next_check TEXT);
CREATE TABLE IF NOT EXISTS bookings(
  id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER, qty INTEGER,
  event TEXT, start TEXT, end TEXT, person TEXT);
CREATE TABLE IF NOT EXISTS checks(
  id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER, day TEXT,
  status TEXT, note TEXT);
CREATE TABLE IF NOT EXISTS purchases(
  id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER, name TEXT, qty INTEGER,
  reason TEXT, status TEXT DEFAULT 'Запланирована', created TEXT);

CREATE TABLE IF NOT EXISTS users(
  id INTEGER PRIMARY KEY AUTOINCREMENT, login TEXT UNIQUE NOT NULL COLLATE NOCASE,
  pw_hash TEXT NOT NULL, full_name TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'user',
  employee_id INTEGER, active INTEGER NOT NULL DEFAULT 1, created TEXT, last_login TEXT);
CREATE TABLE IF NOT EXISTS categories(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE NOT NULL);
CREATE TABLE IF NOT EXISTS states(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE NOT NULL);
CREATE TABLE IF NOT EXISTS locations(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE NOT NULL);
CREATE TABLE IF NOT EXISTS employees(
  id INTEGER PRIMARY KEY AUTOINCREMENT, full_name TEXT NOT NULL, position TEXT, phone TEXT,
  email TEXT, active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS events(
  id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, event_date TEXT, place TEXT, note TEXT);
CREATE TABLE IF NOT EXISTS issues(
  id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER NOT NULL, employee_id INTEGER NOT NULL,
  event_id INTEGER, qty INTEGER NOT NULL, issued_at TEXT NOT NULL, due_date TEXT NOT NULL,
  issued_by INTEGER, issue_note TEXT, returned_at TEXT, return_state TEXT, return_note TEXT,
  received_by INTEGER, return_photos TEXT);
CREATE TABLE IF NOT EXISTS movements(
  id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER NOT NULL, ts TEXT NOT NULL, kind TEXT NOT NULL,
  qty INTEGER, employee_id INTEGER, event_id INTEGER, user_id INTEGER, details TEXT);
CREATE TABLE IF NOT EXISTS audit_log(
  id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, user_id INTEGER, login TEXT,
  action TEXT NOT NULL, entity TEXT, entity_id TEXT, details TEXT, ip TEXT);
CREATE TABLE IF NOT EXISTS inventories(
  id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, created TEXT, finished TEXT,
  status TEXT DEFAULT 'Идёт', scope TEXT, created_by INTEGER, applied INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS inventory_lines(
  id INTEGER PRIMARY KEY AUTOINCREMENT, inv_id INTEGER NOT NULL, item_id INTEGER NOT NULL,
  expected INTEGER NOT NULL, issued INTEGER DEFAULT 0, actual INTEGER, note TEXT);
CREATE INDEX IF NOT EXISTS ix_issues_item ON issues(item_id);
CREATE INDEX IF NOT EXISTS ix_issues_open ON issues(returned_at, due_date);
CREATE INDEX IF NOT EXISTS ix_mov_item ON movements(item_id);
CREATE INDEX IF NOT EXISTS ix_audit_ts ON audit_log(ts);
CREATE INDEX IF NOT EXISTS ix_invl_inv ON inventory_lines(inv_id);
"""
# новые столбцы для баз, созданных прежней версией приложения
MIGRATIONS = [("items", "descr", "TEXT"), ("items", "serial", "TEXT"),
              ("bookings", "event_id", "INTEGER"), ("bookings", "employee_id", "INTEGER"),
              ("issues", "return_photos", "TEXT")]


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB, timeout=15)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d:
        d.close()


def q(sql, args=(), one=False):
    cur = db().execute(sql, args)
    rows = [dict(r) for r in cur.fetchall()]
    return (rows[0] if rows else None) if one else rows


def x(sql, args=()):
    """Выполнить без commit (для составных операций; завершать db().commit())."""
    return db().execute(sql, args)


def run(sql, args=()):
    db().execute(sql, args)
    db().commit()


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def today():
    return date.today().isoformat()


def init_db():
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    con = sqlite3.connect(DB, timeout=15)
    con.executescript(SCHEMA)
    for table, col, ddl in MIGRATIONS:
        cols = [r[1] for r in con.execute(f"PRAGMA table_info({table})")]
        if col not in cols:
            con.execute(f"ALTER TABLE {table} ADD COLUMN {col} {ddl}")

    # справочники: стартовые значения + всё, что уже используется в карточках
    for s in SYSTEM_STATES:
        con.execute("INSERT OR IGNORE INTO states(name) VALUES(?)", (s,))
    for tbl, col, defaults in (("categories", "category", DEFAULT_CATEGORIES),
                               ("locations", "location", DEFAULT_LOCATIONS),
                               ("states", "status", [])):
        if not con.execute(f"SELECT 1 FROM {tbl}").fetchone() or tbl == "states":
            for n in defaults:
                con.execute(f"INSERT OR IGNORE INTO {tbl}(name) VALUES(?)", (n,))
        for (n,) in con.execute(f"SELECT DISTINCT {col} FROM items WHERE {col} IS NOT NULL AND {col}<>''").fetchall():
            con.execute(f"INSERT OR IGNORE INTO {tbl}(name) VALUES(?)", (n,))
    if not con.execute("SELECT 1 FROM events").fetchone():
        con.execute("INSERT INTO events(name, note) VALUES(?,?)",
                    (NO_EVENT, "Служебная запись для операций вне мероприятий"))

    fresh_demo = False
    if not con.execute("SELECT 1 FROM items").fetchone():
        fresh_demo = True
        _seed_demo(con)
    if not con.execute("SELECT 1 FROM users").fetchone():
        con.execute("INSERT INTO users(login,pw_hash,full_name,role,created) VALUES(?,?,?,?,?)",
                    (ADMIN_LOGIN, generate_password_hash(ADMIN_PASS), "Администратор", "admin", now()))
        if fresh_demo:  # демонстрационные учётные записи только для новой демо-базы
            emp = con.execute("SELECT id FROM employees ORDER BY id").fetchall()
            con.execute("INSERT INTO users(login,pw_hash,full_name,role,employee_id,created) VALUES(?,?,?,?,?,?)",
                        ("manager", generate_password_hash("manager"), "Кузнецова А. П.", "responsible",
                         emp[0][0] if emp else None, now()))
            con.execute("INSERT INTO users(login,pw_hash,full_name,role,employee_id,created) VALUES(?,?,?,?,?,?)",
                        ("user", generate_password_hash("user"), "Иванов С. В.", "user",
                         emp[1][0] if len(emp) > 1 else None, now()))
    con.commit()
    con.close()


def _seed_demo(con):
    t = date.today()
    demo = [
        ("АГФ-001", "Рояль концертный Steinway", "Музыкальные инструменты", "Большой зал", 1, 1, "Исправно", 6500000, 12, 30),
        ("АГФ-002", "Микрофон Shure SM58", "Звук", "Склад звука", 12, 8, "Исправно", 9500, 6, 40),
        ("АГФ-003", "Пульт музыкантский", "Сценическое оборудование", "Склад сцены", 40, 30, "Исправно", 3200, 10, 90),
        ("АГФ-004", "Прожектор LED PAR 64", "Свет", "Склад света", 6, 10, "Требует ремонта", 14000, 5, -10),
        ("АГФ-005", "Микшерный пульт Yamaha", "Звук", "Аппаратная", 1, 1, "Исправно", 210000, 7, -5),
        ("АГФ-006", "Театральный реквизит: кресло XIX в.", "Реквизит", "Склад реквизита", 2, 2, "Исправно", 25000, 15, 120),
    ]
    ids = []
    for i, (inv, n, c, loc, qty, mq, st, pr, life, chk) in enumerate(demo):
        bought = (t - timedelta(days=365 * (life - 0.5 if i == 3 else 2))).isoformat()
        cur = con.execute("INSERT INTO items(inv_no,name,category,location,qty,min_qty,status,price,"
                          "bought,life_years,next_check) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                          (inv, n, c, loc, qty, mq, st, pr, bought, life, (t + timedelta(days=chk)).isoformat()))
        ids.append(cur.lastrowid)
        con.execute("INSERT INTO movements(item_id,ts,kind,qty,details) VALUES(?,?,?,?,?)",
                    (cur.lastrowid, now(), "Постановка на учет", qty, f"Инв. № {inv} (начальные данные)"))
    emps = [("Кузнецова Анна Петровна", "Заведующая реквизитом", "+7 (8512) 00-00-01", "kuznetsova@example.org"),
            ("Иванов Сергей Викторович", "Звукорежиссёр", "+7 (8512) 00-00-02", "ivanov@example.org"),
            ("Сидоров Павел Андреевич", "Художник по свету", "+7 (8512) 00-00-03", "sidorov@example.org")]
    eids = [con.execute("INSERT INTO employees(full_name,position,phone,email) VALUES(?,?,?,?)", e).lastrowid
            for e in emps]
    evs = [("Концерт симфонического оркестра", (t - timedelta(days=3)).isoformat(), "Большой зал"),
           ("Гала-концерт «Астраханская осень»", (t + timedelta(days=5)).isoformat(), "Большой зал")]
    evids = [con.execute("INSERT INTO events(name,event_date,place) VALUES(?,?,?)", e).lastrowid for e in evs]
    # демонстрация уведомлений: одна просроченная выдача и одна со скорым сроком
    for item_i, emp, ev, qty, out_days, due_days in ((1, eids[1], evids[0], 2, 5, -1), (2, eids[0], evids[1], 10, 1, 2)):
        iid = ids[item_i]
        issued_at, due = (t - timedelta(days=out_days)).isoformat(), (t + timedelta(days=due_days)).isoformat()
        con.execute("INSERT INTO issues(item_id,employee_id,event_id,qty,issued_at,due_date,issue_note) "
                    "VALUES(?,?,?,?,?,?,?)", (iid, emp, ev, qty, issued_at, due, "Демонстрационная выдача"))
        con.execute("INSERT INTO movements(item_id,ts,kind,qty,employee_id,event_id,details) VALUES(?,?,?,?,?,?,?)",
                    (iid, issued_at + " 10:00:00", "Выдача", qty, emp, ev, f"Срок возврата {due}"))


# ---------------------------------------------------------------- вспомогательное
def to_int(v, default=None):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return default


def to_float(v, default=None):
    try:
        r = float(str(v).replace(",", ".").strip())
        return r if math.isfinite(r) else default
    except (TypeError, ValueError):
        return default


def valid_date(v):
    try:
        date.fromisoformat(v or "")
        return True
    except ValueError:
        return False


def names(table):
    return [r["name"] for r in q(f"SELECT name FROM {table} ORDER BY name")]


def active_employees():
    return q("SELECT id, full_name, position FROM employees WHERE active=1 ORDER BY full_name")


def event_choices():
    return q("SELECT id, name, event_date FROM events ORDER BY COALESCE(event_date,'9999') DESC, name")


def stage(item):
    """Этап жизненного цикла."""
    if item["status"] == "Списано":
        return "Списание"
    if item["status"] in ("Требует ремонта", "В ремонте"):
        return "Ремонт"
    return "Эксплуатация"


def avail_status(it):
    """Статус объекта (наличие), вычисляется автоматически."""
    issued = it.get("issued", 0) or 0
    if it["status"] == "Списано":
        return "Списан"
    if it["status"] == "В ремонте":
        return "В ремонте"
    if it["qty"] <= 0:
        return "Отсутствует"
    if issued >= it["qty"]:
        return "Выдан"
    if issued > 0:
        return "Выдан частично"
    return "В наличии"


# ---------------------------------------------------------------- 3D-модели инструментов
INSTR_CATEGORY_HINT = "инструмент"  # категории, в названии которых есть это слово, считаются музыкальными
INSTR3D_TITLES = {
    "grand": "рояль", "upright": "пианино", "synth": "синтезатор", "violin": "скрипка", "viola": "альт",
    "cello": "виолончель", "bass": "контрабас", "guitar": "гитара", "electric": "электрогитара",
    "bassg": "бас-гитара", "balalaika": "балалайка", "domra": "домра", "mandolin": "мандолина", "harp": "арфа",
    "flute": "флейта", "oboe": "гобой", "clarinet": "кларнет", "bassoon": "фагот", "sax": "саксофон",
    "trumpet": "труба", "trombone": "тромбон", "horn": "валторна", "tuba": "туба", "snare": "малый барабан",
    "bassdrum": "большой барабан", "timpani": "литавры", "cymbals": "тарелки", "xylo": "ксилофон / маримба",
    "accordion": "аккордеон / баян", "organ": "орган", "note": "музыкальный инструмент"}
# порядок важен: первое совпадение побеждает (синтезатор раньше пианино, бас-гитара раньше гитары и т. д.)
INSTR3D_RULES = [(k, re.compile(p)) for k, p in (
    ("synth", r"синтезатор|цифров\w*\s+(?:пианино|фортепиано|рояль)|электропиано|клавишн"),
    ("grand", r"роял"),
    ("upright", r"пианин|фортепиано"),
    ("bassg", r"бас[\s-]*гитар|гитар\w*[\s-]+бас"),
    ("electric", r"электро[\s-]*гитар"),
    ("guitar", r"гитар"),
    ("balalaika", r"балалайк"),
    ("domra", r"домр"),
    ("mandolin", r"мандолин"),
    ("violin", r"скрипк"),
    ("cello", r"виолончел"),
    ("bass", r"контрабас"),
    ("sax", r"саксофон"),
    ("viola", r"(?<![а-я])альт(?![а-я])|(?<![а-я])виол(?:а|ы|у|е|ой)(?![а-я])"),
    ("harp", r"арф"),
    ("flute", r"флейт|пикколо"),
    ("oboe", r"гобо"),
    ("clarinet", r"кларнет"),
    ("bassoon", r"фагот"),
    ("trombone", r"тромбон"),
    ("horn", r"валторн|(?<![а-я])горн"),
    ("tuba", r"(?<![а-я])туб(?![а-я]*ус)|геликон|эуфониум|басов\w*\s+труб"),
    ("trumpet", r"труб(?!ч)|корнет|флюгельгорн"),
    ("bassdrum", r"(?:бас|больш)\w*[\s-]+барабан|барабан\w*[\s-]+(?:бас|больш)"),
    ("timpani", r"литавр|тимпан"),
    ("snare", r"барабан"),
    ("cymbals", r"тарелк"),
    ("xylo", r"ксилофон|маримб|вибрафон|металлофон|глокеншпиль"),
    ("accordion", r"аккордеон|баян|гармон"),
    ("organ", r"(?<![а-я])орган(?![а-я])"),
)]
# принадлежности («чехол для скрипки», «струны», «стойка»…) — это не инструмент, модель не показываем
INSTR3D_ACCESSORY = re.compile(r"^\s*(?:чехол|кейс|футляр|струн|смыч|канифол|медиатор|палоч|мундштук|трост|сурдин|ремен|ремн|"
                               r"подставк|стойк|пюпитр|пульт|тюнер|метроном|педал|крышк|накладк|кожа|набор|комплект|запчаст|аксессуар)")


def instrument_kind(name, category=""):
    """Тип 3D-модели по названию карточки (None — модель не нужна). Музыкальная категория без опознанного
    названия получает универсальную модель «нота»."""
    n = (name or "").lower().replace("ё", "е")
    if INSTR3D_ACCESSORY.match(n) or re.search(r"\s(?:для|под)\s", n):
        return None
    for kind, rx in INSTR3D_RULES:
        if rx.search(n):
            return kind
    if INSTR_CATEGORY_HINT in (category or "").lower():
        return "note"
    return None



def enrich(rows):
    for r in rows:
        r["wear"] = wear(r)
        r["wear_level"] = wear_level(r["wear"])
        r["wear_description"] = wear_description(r["wear"], r.get("bought"), r.get("life_years"))
        r["stage"] = stage(r)
        r["overdue"] = bool(r["next_check"] and r["next_check"] < date.today().isoformat()
                            and r["status"] != "Списано")
        r["issued"] = r.get("issued") or 0
        r["free"] = r["qty"] - r["issued"]
        r["avail"] = avail_status(r)
        r["m3d"] = instrument_kind(r.get("name"), r.get("category"))
        r["m3d_title"] = INSTR3D_TITLES.get(r["m3d"], "")
    return rows


ITEM_SQL = ("SELECT i.*, COALESCE((SELECT SUM(s.qty) FROM issues s WHERE s.item_id=i.id "
            "AND s.returned_at IS NULL),0) AS issued FROM items i")


def get_item(iid):
    r = q(ITEM_SQL + " WHERE i.id=?", (iid,), one=True)
    return enrich([r])[0] if r else None


def next_inv_no():
    mx = 0
    pat = re.compile(re.escape(INV_PREFIX) + r"(\d+)")
    for r in q("SELECT inv_no FROM items"):
        m = pat.fullmatch(r["inv_no"] or "")
        if m:
            mx = max(mx, int(m.group(1)))
    return f"{INV_PREFIX}{mx + 1:03d}"


def free_qty(item_id, start, end):
    """Свободно на период: всего − брони − невозвращённые выдачи, пересекающиеся с периодом."""
    item = q("SELECT qty FROM items WHERE id=?", (item_id,), one=True)
    booked = q("SELECT COALESCE(SUM(qty),0) s FROM bookings WHERE item_id=? AND start<=? AND end>=?",
               (item_id, end, start), one=True)["s"]
    out = q("SELECT COALESCE(SUM(qty),0) s FROM issues WHERE item_id=? AND returned_at IS NULL "
            "AND issued_at<=? AND MAX(due_date,?)>=?", (item_id, end, today(), start), one=True)["s"]
    return item["qty"] - booked - out


def suggest_purchases():
    """Автоматическое формирование плана закупок."""
    added = 0
    for it in enrich(q(ITEM_SQL)):
        reason = None
        if it["status"] == "Списано":
            reason = "Объект списан — требуется замена"
        elif it["wear"] >= 100:
            reason = "Исчерпан срок службы"
        elif it["qty"] < it["min_qty"]:
            reason = f"Остаток {it['qty']} ниже минимального {it['min_qty']}"
        elif it["status"] == "Требует ремонта" and it["wear"] >= 80:
            reason = "Ремонт нецелесообразен (износ ≥ 80%)"
        if reason and not q("SELECT 1 FROM purchases WHERE item_id=? AND status IN "
                            "('Запланирована','Заказана')", (it["id"],), one=True):
            need = max(1, it["min_qty"] - it["qty"]) if it["qty"] < it["min_qty"] else 1
            x("INSERT INTO purchases(item_id,name,qty,reason,created) VALUES(?,?,?,?,?)",
              (it["id"], it["name"], need, reason, date.today().isoformat()))
            added += 1
    db().commit()
    return added


# ---------------------------------------------------------------- журнал и история
def write_audit(con, user_id, login, action, entity="", entity_id="", details="", ip=""):
    con.execute("INSERT INTO audit_log(ts,user_id,login,action,entity,entity_id,details,ip) "
                "VALUES(?,?,?,?,?,?,?,?)", (now(), user_id, login, action, entity, str(entity_id), details, ip))


def audit(action, entity="", entity_id="", details=""):
    """Запись в журнал действий (commit — на вызывающей стороне)."""
    u = g.get("user") if has_request_context() else None
    ip = request.remote_addr if has_request_context() else ""
    write_audit(db(), u["id"] if u else None, u["login"] if u else None, action, entity, entity_id, details, ip)


def log_move(item_id, kind, qty=None, details="", employee_id=None, event_id=None):
    """Запись в историю движения объекта (commit — на вызывающей стороне)."""
    u = g.get("user")
    x("INSERT INTO movements(item_id,ts,kind,qty,employee_id,event_id,user_id,details) VALUES(?,?,?,?,?,?,?,?)",
      (item_id, now(), kind, qty, employee_id, event_id, u["id"] if u else None, details))


# ---------------------------------------------------------------- аутентификация и права
FAILS = {}  # (ip, login) -> [время неудачных попыток]


def throttled(key):
    t = time.time()
    FAILS[key] = [v for v in FAILS.get(key, []) if t - v < 300]
    return len(FAILS[key]) >= 5


def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]


def csrf_input():
    return Markup('<input type="hidden" name="_csrf" value="%s">' % csrf_token())


@app.before_request
def load_user():
    g.user = None
    uid = session.get("uid")
    if uid:
        u = q("SELECT * FROM users WHERE id=? AND active=1", (uid,), one=True)
        if u:
            g.user = u
        else:
            session.clear()
    if request.method == "POST":
        sent, real = request.form.get("_csrf", ""), session.get("csrf", "")
        if not real or not hmac.compare_digest(sent.encode(), real.encode()):
            abort(400)


def can(perm):
    return bool(g.get("user")) and g.user["role"] in PERMS[perm]


def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if not g.user:
            return redirect(url_for("login", next=request.full_path.rstrip("?")))
        return f(*a, **k)
    return w


def need(perm):
    def deco(f):
        @wraps(f)
        def w(*a, **k):
            if not g.user:
                return redirect(url_for("login", next=request.full_path.rstrip("?")))
            if not can(perm):
                abort(403)
            return f(*a, **k)
        return w
    return deco


def scope_employee():
    """Для обычного пользователя — id связанного сотрудника (видит только свои выдачи); для staff — None."""
    if can("operations"):
        return None
    return g.user["employee_id"] or -1


def validate_password(p):
    return None if len(p or "") >= 6 else "Пароль должен содержать не менее 6 символов"


@app.errorhandler(403)
def e403(_):
    return render_template("error.html", code=403, text="Недостаточно прав для этого действия"), 403


@app.errorhandler(404)
def e404(_):
    return render_template("error.html", code=404, text="Страница не найдена"), 404


@app.errorhandler(400)
def e400(_):
    return render_template("error.html", code=400,
                           text="Сессия устарела или запрос некорректен. Обновите страницу и повторите действие"), 400


# ---------------------------------------------------------------- вход / профиль
def safe_next(n):
    return n if n and n.startswith("/") and not n.startswith("//") and "\\" not in n else url_for("dashboard")


@app.route("/assets/instrument3d.js")
def instrument3d_js():
    """Автономный WebGL-просмотрщик 3D-моделей (без внешних библиотек и интернета)."""
    resp = app.response_class(INSTRUMENT3D_JS, mimetype="application/javascript")
    resp.headers["Cache-Control"] = "public, max-age=3600"
    return resp


@app.route("/login", methods=["GET", "POST"])
def login():
    if g.user:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        lg, pw = request.form.get("login", "").strip(), request.form.get("password", "")
        key = (request.remote_addr, lg.lower())
        if throttled(key):
            audit("Вход заблокирован (слишком много попыток)", "users", "", f"логин: {lg}")
            db().commit()
            flash("Слишком много неудачных попыток. Повторите через несколько минут", "danger")
            return render_template("login.html"), 429
        u = q("SELECT * FROM users WHERE login=?", (lg,), one=True)
        ok = check_password_hash(u["pw_hash"] if u else generate_password_hash("x"), pw)
        if u and ok and u["active"]:
            session.clear()
            session["uid"] = u["id"]
            session.permanent = True
            FAILS.pop(key, None)
            g.user = u
            x("UPDATE users SET last_login=? WHERE id=?", (now(), u["id"]))
            audit("Вход в систему", "users", u["id"])
            db().commit()
            return redirect(safe_next(request.args.get("next")))
        FAILS.setdefault(key, []).append(time.time())
        audit("Неудачная попытка входа", "users", "", f"логин: {lg}")
        db().commit()
        flash("Неверный логин или пароль", "danger")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    if g.user:
        audit("Выход из системы", "users", g.user["id"])
        db().commit()
    session.clear()
    return redirect(url_for("login"))


@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    if request.method == "POST":
        f = request.form
        err = None
        if not check_password_hash(g.user["pw_hash"], f.get("old", "")):
            err = "Текущий пароль указан неверно"
        elif f.get("new") != f.get("new2"):
            err = "Новый пароль и подтверждение не совпадают"
        else:
            err = validate_password(f.get("new"))
        if err:
            flash(err, "danger")
        else:
            x("UPDATE users SET pw_hash=? WHERE id=?", (generate_password_hash(f["new"]), g.user["id"]))
            audit("Смена пароля", "users", g.user["id"])
            db().commit()
            flash("Пароль изменён", "success")
            return redirect(url_for("profile"))
    emp = q("SELECT * FROM employees WHERE id=?", (g.user["employee_id"],), one=True) if g.user["employee_id"] else None
    return render_template("profile.html", emp=emp)


# ---------------------------------------------------------------- уведомления
ISSUE_SQL = ("SELECT s.*, i.name AS item_name, i.inv_no, e.full_name AS emp_name, e.email AS emp_email, "
             "ev.name AS event_name, u.full_name AS issued_by_name, ur.full_name AS received_by_name "
             "FROM issues s JOIN items i ON i.id=s.item_id JOIN employees e ON e.id=s.employee_id "
             "LEFT JOIN events ev ON ev.id=s.event_id LEFT JOIN users u ON u.id=s.issued_by "
             "LEFT JOIN users ur ON ur.id=s.received_by")


def decorate_issue(r):
    if r["returned_at"]:
        r["days"], r["flag"] = None, "returned"
    else:
        r["days"] = (date.fromisoformat(r["due_date"]) - date.today()).days
        r["flag"] = "overdue" if r["days"] < 0 else "soon" if r["days"] <= REMIND_DAYS else "ok"
    return r


def issue_alerts(employee_id=None):
    """Открытые выдачи с просроченным или скорым сроком возврата."""
    limit = (date.today() + timedelta(days=REMIND_DAYS)).isoformat()
    sql, args = ISSUE_SQL + " WHERE s.returned_at IS NULL AND s.due_date<=?", [limit]
    if employee_id is not None:
        sql += " AND s.employee_id=?"
        args.append(employee_id)
    return [decorate_issue(r) for r in q(sql + " ORDER BY s.due_date, s.id", args)]


@app.context_processor
def inject():
    n = 0
    if g.get("user"):
        n = len(issue_alerts(scope_employee()))
        if can("operations"):
            n += q("SELECT COUNT(*) c FROM items WHERE next_check<? AND status<>'Списано'", (today(),), one=True)["c"]
    return dict(user=g.get("user"), can=can, csrf_input=csrf_input, ROLES=ROLES, today=today(),
                notif_count=n, REMIND_DAYS=REMIND_DAYS)


@app.route("/notifications")
@login_required
def notifications():
    alerts = issue_alerts(scope_employee())
    staff = can("operations")
    items = enrich(q(ITEM_SQL)) if staff else []
    return render_template("notifications.html",
                           overdue=[a for a in alerts if a["flag"] == "overdue"],
                           soon=[a for a in alerts if a["flag"] != "overdue"],
                           checks=[i for i in items if i["overdue"]],
                           low=[i for i in items if i["qty"] < i["min_qty"] and i["status"] != "Списано"],
                           smtp_on=bool(os.environ.get("SMTP_HOST")))


def send_reminders():
    """E-mail-напоминания сотрудникам. Возвращает (отправлено, без e-mail, ошибка)."""
    host = os.environ.get("SMTP_HOST")
    if not host:
        return 0, 0, "SMTP не настроен (переменная окружения SMTP_HOST)"
    by_emp = {}
    for a in issue_alerts():
        by_emp.setdefault((a["employee_id"], a["emp_name"], a["emp_email"]), []).append(a)
    sent = skipped = 0
    try:
        srv = smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", "25")), timeout=20)
        if os.environ.get("SMTP_TLS", "0") == "1":
            srv.starttls()
        if os.environ.get("SMTP_USER"):
            srv.login(os.environ["SMTP_USER"], os.environ.get("SMTP_PASSWORD", ""))
        for (_, name, email), rows in by_emp.items():
            if not email:
                skipped += 1
                continue
            lines = []
            for a in rows:
                state = (f"ПРОСРОЧЕНО на {-a['days']} дн." if a["days"] < 0
                         else "срок возврата сегодня" if a["days"] == 0 else f"осталось дней: {a['days']}")
                lines.append(f"— {a['item_name']} ({a['inv_no']}), {a['qty']} шт., вернуть до {a['due_date']} — {state}")
            msg = EmailMessage()
            msg["Subject"] = "Напоминание о возврате имущества"
            msg["From"] = os.environ.get("SMTP_FROM", "inventory@localhost")
            msg["To"] = email
            msg.set_content(f"Здравствуйте, {name}!\n\nСроки возврата закреплённого за вами имущества:\n\n"
                            + "\n".join(lines) + "\n\nАвтоматическое сообщение системы учета.")
            srv.send_message(msg)
            sent += 1
        srv.quit()
    except (OSError, smtplib.SMTPException) as e:
        return sent, skipped, f"Ошибка отправки: {e}"
    return sent, skipped, None


@app.route("/notifications/send", methods=["POST"])
@need("operations")
def notifications_send():
    sent, skipped, err = send_reminders()
    audit("Отправка напоминаний", "notifications", "", f"отправлено: {sent}, без e-mail: {skipped}" + (f"; {err}" if err else ""))
    db().commit()
    if err:
        flash(err, "danger")
    else:
        flash(f"Отправлено писем: {sent}" + (f"; у {skipped} сотрудников не указан e-mail" if skipped else ""), "success")
    return redirect(url_for("notifications"))


# ---------------------------------------------------------------- главная
@app.route("/")
@login_required
def dashboard():
    items = enrich(q(ITEM_SQL))
    emp = scope_employee()
    alerts = issue_alerts(emp)
    stats = {
        "total": len(items),
        "ok": sum(i["status"] == "Исправно" for i in items),
        "repair": sum(i["stage"] == "Ремонт" for i in items),
        "overdue": [i for i in items if i["overdue"]],
        "worn": [i for i in items if i["wear"] >= 80 and i["status"] != "Списано"],
        "value": sum(i["price"] * i["qty"] for i in items if i["status"] != "Списано"),
        "issued": sum(i["issued"] for i in items),
        "late": [a for a in alerts if a["flag"] == "overdue"],
        "soon": [a for a in alerts if a["flag"] != "overdue"],
    }
    upcoming = q("SELECT b.*, i.name FROM bookings b JOIN items i ON i.id=b.item_id "
                 "WHERE b.end>=? ORDER BY b.start LIMIT 8", (today(),))
    return render_template("dashboard.html", s=stats, upcoming=upcoming)


# ---------------------------------------------------------------- имущество
FIELD_LABELS = {"inv_no": "Инв. №", "name": "Название", "category": "Категория", "location": "Место хранения",
                "qty": "Количество", "min_qty": "Мин. запас", "price": "Цена", "bought": "Дата покупки",
                "life_years": "Срок службы", "descr": "Описание", "serial": "Серийный №"}


def search_items(a):
    """Поиск и фильтрация: текст, категория, место, статус (наличие), состояние."""
    s = (a.get("s") or "").strip().casefold()
    sql, args, w = ITEM_SQL, [], []
    for key, col in (("cat", "category"), ("loc", "location"), ("state", "status")):
        if a.get(key):
            w.append(f"i.{col}=?")
            args.append(a[key])
    rows = enrich(q(sql + (" WHERE " + " AND ".join(w) if w else "") + " ORDER BY i.name", args))
    if a.get("status"):
        rows = [r for r in rows if r["avail"] == a["status"]]
    if s:
        rows = [r for r in rows if s in " ".join(str(r.get(k) or "") for k in ("name", "inv_no", "serial", "descr")).casefold()]
    return rows


def item_filters_ctx():
    return dict(cats=names("categories"), locs=names("locations"), states=names("states"), avail=AVAIL_STATUSES)


@app.route("/items")
@login_required
def items():
    a = request.args
    return render_template("items.html", items=search_items(a), a=a, **item_filters_ctx())


def parse_item(f, min_total=1, issued=0):
    """Проверка формы карточки. Возвращает (данные, ошибки)."""
    err, d = [], {}
    d["inv_no"] = f.get("inv_no", "").strip()
    d["name"] = f.get("name", "").strip()
    d["descr"] = f.get("descr", "").strip()
    d["serial"] = f.get("serial", "").strip()
    if not d["name"]:
        err.append("Укажите название")
    d["category"] = f.get("category", "").strip()
    if d["category"] not in names("categories"):
        err.append("Выберите категорию из справочника")
    d["location"] = f.get("location", "").strip()
    if d["location"] not in names("locations"):
        err.append("Выберите место хранения из справочника")
    d["qty"] = to_int(f.get("qty"))
    if d["qty"] is None or d["qty"] < min_total:
        err.append(f"Количество должно быть целым числом не меньше {min_total}"
                   + (f" (сейчас выдано {issued} шт.)" if issued and min_total == issued else ""))
    d["min_qty"] = to_int(f.get("min_qty"), 0)
    if d["min_qty"] < 0:
        err.append("Минимальный запас не может быть отрицательным")
    d["price"] = to_float(f.get("price"), None)
    if d["price"] is None or d["price"] < 0:
        err.append("Цена указана некорректно")
    d["bought"] = f.get("bought", "").strip() or today()
    if not valid_date(d["bought"]):
        err.append("Дата покупки указана некорректно")
    d["life_years"] = to_float(f.get("life_years"), None)
    if d["life_years"] is None or d["life_years"] <= 0:
        err.append("Срок службы должен быть положительным числом")
    return d, err


@app.route("/items/new", methods=["GET", "POST"])
@need("edit_items")
def item_new():
    form = request.form if request.method == "POST" else {"qty": 1, "min_qty": 0, "price": 0, "life_years": 5,
                                                         "bought": today(), "state": OK_STATE}
    if request.method == "POST":
        d, err = parse_item(form, 1)
        state = form.get("state", OK_STATE)
        if state not in names("states"):
            err.append("Выберите состояние из справочника")
        if not err:
            auto = not d["inv_no"]
            for _ in range(5 if auto else 1):
                inv = d["inv_no"] or next_inv_no()
                try:
                    cur = x("INSERT INTO items(inv_no,name,category,location,qty,min_qty,status,price,bought,"
                            "life_years,next_check,descr,serial) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                            (inv, d["name"], d["category"], d["location"], d["qty"], d["min_qty"], state,
                             d["price"], d["bought"], d["life_years"],
                             (date.today() + timedelta(days=CHECK_PERIOD_DAYS)).isoformat(), d["descr"], d["serial"]))
                    break
                except sqlite3.IntegrityError:
                    inv = None
            if inv is None:
                err.append("Инвентарный номер уже существует")
            else:
                log_move(cur.lastrowid, "Постановка на учет", d["qty"],
                         f"Инв. № {inv}" + (" (присвоен автоматически)" if auto else "") + f"; состояние: {state}")
                audit("Добавлена карточка", "items", cur.lastrowid, f"{inv} {d['name']}, {d['qty']} шт.")
                db().commit()
                flash(f"Объект поставлен на учет. Инвентарный номер: {inv}", "success")
                return redirect(url_for("item", iid=cur.lastrowid))
        for e in err:
            flash(e, "danger")
    return render_template("item_form.html", it=None, f=form, next_inv=next_inv_no(), **item_filters_ctx())


@app.route("/items/<int:iid>/edit", methods=["GET", "POST"])
@need("edit_items")
def item_edit(iid):
    it = get_item(iid)
    if not it:
        abort(404)
    form = request.form if request.method == "POST" else it
    if request.method == "POST":
        d, err = parse_item(form, max(it["issued"], 0), it["issued"])
        if can("admin") and d["inv_no"]:
            pass
        else:
            d["inv_no"] = it["inv_no"]  # менять инв. номер может только администратор
        if not err:
            try:
                x("UPDATE items SET inv_no=?,name=?,category=?,location=?,qty=?,min_qty=?,price=?,bought=?,"
                  "life_years=?,descr=?,serial=? WHERE id=?",
                  (d["inv_no"], d["name"], d["category"], d["location"], d["qty"], d["min_qty"], d["price"],
                   d["bought"], d["life_years"], d["descr"], d["serial"], iid))
            except sqlite3.IntegrityError:
                err.append("Инвентарный номер уже существует")
        if not err:
            diff = [f"{FIELD_LABELS[k]}: {it.get(k) or '—'} → {d[k] or '—'}" for k in FIELD_LABELS
                    if str(it.get(k) if it.get(k) is not None else "") != str(d[k] if d[k] is not None else "")]
            if diff:
                kind = "Перемещение" if it["location"] != d["location"] else "Изменение карточки"
                log_move(iid, kind, d["qty"], "; ".join(diff))
                audit("Изменена карточка", "items", iid, "; ".join(diff))
            db().commit()
            flash("Карточка сохранена", "success")
            return redirect(url_for("item", iid=iid))
        db().rollback()
        for e in err:
            flash(e, "danger")
    return render_template("item_form.html", it=it, f=form, next_inv=None, **item_filters_ctx())


MOVE_SQL = ("SELECT m.*, i.name AS item_name, i.inv_no, e.full_name AS emp_name, ev.name AS event_name, "
            "u.full_name AS user_name FROM movements m JOIN items i ON i.id=m.item_id "
            "LEFT JOIN employees e ON e.id=m.employee_id LEFT JOIN events ev ON ev.id=m.event_id "
            "LEFT JOIN users u ON u.id=m.user_id")


@app.route("/items/<int:iid>")
@login_required
def item(iid):
    it = get_item(iid)
    if not it:
        return redirect(url_for("items"))
    staff = can("operations")
    emp = scope_employee()
    issues = [decorate_issue(r) for r in q(ISSUE_SQL + " WHERE s.item_id=?" + ("" if staff else " AND s.employee_id=?")
                                           + " ORDER BY s.id DESC", (iid,) if staff else (iid, emp))]
    return render_template("item.html", it=it, states=names("states"), issues=issues,
                           moves=q(MOVE_SQL + " WHERE m.item_id=? ORDER BY m.id DESC", (iid,)) if staff else [],
                           checks=q("SELECT * FROM checks WHERE item_id=? ORDER BY id DESC", (iid,)),
                           bookings=q("SELECT * FROM bookings WHERE item_id=? ORDER BY start DESC", (iid,)))


@app.route("/items/<int:iid>/check", methods=["POST"])
@need("edit_items")
def item_check(iid):
    it = get_item(iid)
    if not it:
        abort(404)
    st, note = request.form.get("status", ""), request.form.get("note", "").strip()
    if st not in names("states") or not note:
        flash("Выберите состояние и укажите комментарий", "danger")
    elif st == "Списано" and it["issued"]:
        flash(f"Нельзя списать: не возвращено {it['issued']} шт. Сначала оформите возврат", "danger")
    else:
        x("INSERT INTO checks(item_id,day,status,note) VALUES(?,?,?,?)", (iid, today(), st, note))
        x("UPDATE items SET status=?, next_check=? WHERE id=?",
          (st, (date.today() + timedelta(days=CHECK_PERIOD_DAYS)).isoformat(), iid))
        if st != it["status"]:
            log_move(iid, "Смена состояния", None, f"{it['status']} → {st}. {note}")
        else:
            log_move(iid, "Проверка состояния", None, f"{st}. {note}")
        audit("Акт проверки состояния", "items", iid, f"{it['status']} → {st}: {note}")
        db().commit()
        flash("Состояние обновлено", "success")
    return redirect(url_for("item", iid=iid))


# ---------------------------------------------------------------- выдача и возврат
@app.route("/issues")
@login_required
def issues():
    a = request.args
    tab = a.get("tab", "open")
    w, args = [], []
    emp = scope_employee()
    if emp is not None:
        w.append("s.employee_id=?")
        args.append(emp)
    elif a.get("emp"):
        w.append("s.employee_id=?")
        args.append(to_int(a["emp"], -1))
    if a.get("event"):
        w.append("s.event_id=?")
        args.append(to_int(a["event"], -1))
    if tab in ("open", "overdue"):
        w.append("s.returned_at IS NULL")
    if tab == "overdue":
        w.append("s.due_date<?")
        args.append(today())
    if tab == "closed":
        w.append("s.returned_at IS NOT NULL")
    rows = [decorate_issue(r) for r in q(ISSUE_SQL + (" WHERE " + " AND ".join(w) if w else "")
                                          + " ORDER BY s.returned_at IS NOT NULL, s.due_date, s.id DESC", args)]
    s = (a.get("s") or "").strip().casefold()
    if s:
        rows = [r for r in rows if s in f"{r['item_name']} {r['inv_no']} {r['emp_name']}".casefold()]
    return render_template("issues.html", rows=rows, tab=tab, a=a, emps=q("SELECT id, full_name FROM employees ORDER BY full_name"),
                           events=event_choices())


@app.route("/issues/new", methods=["GET", "POST"])
@need("operations")
def issue_new():
    avail = [i for i in enrich(q(ITEM_SQL + " ORDER BY i.name")) if i["status"] == OK_STATE and i["free"] > 0]
    form = request.form if request.method == "POST" else {
        "item_id": request.args.get("item_id", ""), "event_id": request.args.get("event_id", ""),
        "employee_id": request.args.get("employee_id", ""), "qty": request.args.get("qty", 1),
        "due_date": (date.today() + timedelta(days=DEFAULT_ISSUE_DAYS)).isoformat()}
    if request.method == "POST":
        f, err = request.form, []
        it = next((i for i in avail if str(i["id"]) == f.get("item_id")), None)
        emp = q("SELECT * FROM employees WHERE id=? AND active=1", (to_int(f.get("employee_id"), -1),), one=True)
        ev = q("SELECT * FROM events WHERE id=?", (to_int(f.get("event_id"), -1),), one=True)
        qty, due = to_int(f.get("qty")), f.get("due_date", "")
        if not it:
            err.append("Выберите объект: он должен быть исправен и доступен для выдачи")
        elif qty is None or not 1 <= qty <= it["free"]:
            err.append(f"Количество должно быть от 1 до {it['free']} (свободно сейчас)")
        if not emp:
            err.append("Выберите сотрудника из справочника")
        if not ev:
            err.append("Выберите мероприятие (для хозяйственных нужд есть отдельная запись)")
        if not valid_date(due) or due < today():
            err.append("Срок возврата указан некорректно (не раньше сегодняшней даты)")
        if not err:
            cur = x("INSERT INTO issues(item_id,employee_id,event_id,qty,issued_at,due_date,issued_by,issue_note) "
                    "VALUES(?,?,?,?,?,?,?,?)", (it["id"], emp["id"], ev["id"], qty, today(), due,
                                               g.user["id"], f.get("note", "").strip()))
            log_move(it["id"], "Выдача", qty, f"Срок возврата {due}. {f.get('note', '').strip()}".strip(),
                     emp["id"], ev["id"])
            audit("Выдача имущества", "issues", cur.lastrowid, f"{it['inv_no']} {it['name']} ×{qty} → {emp['full_name']}")
            db().commit()
            flash(f"Выдано: {it['name']} × {qty} — {emp['full_name']}", "success")
            return redirect(url_for("issues"))
        for e in err:
            flash(e, "danger")
    return render_template("issue_new.html", f=form, avail=avail, emps=active_employees(), events=event_choices())


@app.route("/issues/<int:sid>/return", methods=["GET", "POST"])
@need("operations")
def issue_return(sid):
    s = q(ISSUE_SQL + " WHERE s.id=?", (sid,), one=True)
    if not s or s["returned_at"]:
        flash("Выдача не найдена или уже закрыта", "warning")
        return redirect(url_for("issues"))
    decorate_issue(s)
    if request.method == "POST":
        f, err = request.form, []
        rq, state, note = to_int(f.get("qty")), f.get("state", ""), f.get("note", "").strip()
        photos = request.files.getlist("photos")
        photos = [ph for ph in photos if ph and ph.filename]
        if len(photos) > 6:
            err.append("Можно приложить не более 6 фотографий")
        for ph in photos:
            ext = secure_filename(ph.filename).rsplit(".", 1)[-1].lower() if "." in ph.filename else ""
            if ext not in ALLOWED_PHOTO_EXTENSIONS:
                err.append("Фотографии должны быть в формате JPG, PNG или WEBP")
                break
        if rq is None or not 1 <= rq <= s["qty"]:
            err.append(f"Возвращаемое количество — от 1 до {s['qty']}")
        if state not in names("states"):
            err.append("Выберите состояние при возврате")
        elif state != OK_STATE and not note:
            err.append("Если состояние отличается от «Исправно», опишите повреждение в комментарии")
        if not err:
            saved_photos = []
            for ph in photos:
                ext = secure_filename(ph.filename).rsplit(".", 1)[-1].lower()
                filename = secrets.token_hex(16) + "." + ext
                os.makedirs(UPLOAD_DIR, exist_ok=True)
                ph.save(os.path.join(UPLOAD_DIR, filename))
                saved_photos.append(filename)
            photos_json = json.dumps(saved_photos, ensure_ascii=False) if saved_photos else None
            it = get_item(s["item_id"])
            if rq < s["qty"]:  # частичный возврат: закрываем часть записи, остаток остаётся открытым
                x("UPDATE issues SET qty=qty-? WHERE id=?", (rq, sid))
                x("INSERT INTO issues(item_id,employee_id,event_id,qty,issued_at,due_date,issued_by,issue_note,"
                  "returned_at,return_state,return_note,received_by,return_photos) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  (s["item_id"], s["employee_id"], s["event_id"], rq, s["issued_at"], s["due_date"],
                   s["issued_by"], s["issue_note"], today(), state, note, g.user["id"], photos_json))
            else:
                x("UPDATE issues SET returned_at=?, return_state=?, return_note=?, received_by=?, return_photos=? WHERE id=?",
                  (today(), state, note, g.user["id"], photos_json, sid))
            extra = ""
            if state != OK_STATE:
                if rq == it["qty"]:  # вернули всю партию — состояние карточки меняется
                    x("INSERT INTO checks(item_id,day,status,note) VALUES(?,?,?,?)",
                      (it["id"], today(), state, f"Возврат от {s['emp_name']}: {note}"))
                    x("UPDATE items SET status=? WHERE id=?", (state, it["id"]))
                    log_move(it["id"], "Смена состояния", None, f"{it['status']} → {state} (при возврате). {note}",
                             s["employee_id"], s["event_id"])
                else:
                    extra = f" (в карточке {it['qty']} шт.: состояние карточки не менялось)"
                    flash("Состояние зафиксировано в истории, но в карточке с несколькими единицами оно не изменено. "
                          "При необходимости уменьшите количество или заведите отдельную карточку на повреждённые единицы.",
                          "warning")
            log_move(s["item_id"], "Возврат", rq, f"Состояние: {state}. {note}".strip() + extra,
                     s["employee_id"], s["event_id"])
            audit("Возврат имущества", "issues", sid, f"{s['inv_no']} ×{rq}, состояние: {state}")
            db().commit()
            flash("Возврат оформлен", "success")
            return redirect(url_for("issues"))
        for e in err:
            flash(e, "danger")
    try:
        s["photo_list"] = json.loads(s["return_photos"] or "[]")
    except (TypeError, ValueError):
        s["photo_list"] = []
    return render_template("issue_return.html", s=s, states=names("states"), f=request.form)


# ---------------------------------------------------------------- бронирование
@app.route("/bookings", methods=["GET", "POST"])
@login_required
def bookings():
    if request.method == "POST":
        if not can("operations"):
            abort(403)
        f = request.form
        it = q("SELECT * FROM items WHERE id=?", (to_int(f.get("item_id"), -1),), one=True)
        ev = q("SELECT * FROM events WHERE id=?", (to_int(f.get("event_id"), -1),), one=True)
        emp = q("SELECT * FROM employees WHERE id=? AND active=1", (to_int(f.get("employee_id"), -1),), one=True)
        qty = to_int(f.get("qty"))
        if not (it and ev and emp and qty and qty > 0 and valid_date(f.get("start")) and valid_date(f.get("end"))):
            flash("Заполните все поля корректно", "danger")
        elif f["end"] < f["start"]:
            flash("Дата окончания раньше даты начала", "danger")
        elif it["status"] != "Исправно":
            flash(f"Нельзя забронировать: состояние «{it['status']}»", "danger")
        elif free_qty(it["id"], f["start"], f["end"]) < qty:
            flash("Недостаточно свободных единиц на выбранные даты (конфликт бронирования)", "danger")
        else:
            cur = x("INSERT INTO bookings(item_id,qty,event,start,end,person,event_id,employee_id) VALUES(?,?,?,?,?,?,?,?)",
                    (it["id"], qty, ev["name"], f["start"], f["end"], emp["full_name"], ev["id"], emp["id"]))
            audit("Создана бронь", "bookings", cur.lastrowid, f"{it['inv_no']} ×{qty}, {f['start']}…{f['end']}, {ev['name']}")
            db().commit()
            flash("Бронь создана", "success")
        return redirect(url_for("bookings"))
    rows = q("SELECT b.*, i.name, i.inv_no FROM bookings b JOIN items i ON i.id=b.item_id ORDER BY b.start DESC")
    return render_template("bookings.html", rows=rows, events=event_choices(), emps=active_employees(),
                           items=q("SELECT id,name,inv_no,qty FROM items WHERE status='Исправно' ORDER BY name"))


@app.route("/bookings/<int:bid>/delete", methods=["POST"])
@need("operations")
def booking_delete(bid):
    b = q("SELECT b.*, i.inv_no FROM bookings b JOIN items i ON i.id=b.item_id WHERE b.id=?", (bid,), one=True)
    if b:
        x("DELETE FROM bookings WHERE id=?", (bid,))
        audit("Отменена бронь", "bookings", bid, f"{b['inv_no']} ×{b['qty']}, {b['event']}")
        db().commit()
        flash("Бронь отменена", "info")
    return redirect(url_for("bookings"))


# ---------------------------------------------------------------- закупки
@app.route("/purchases", methods=["GET", "POST"])
@need("operations")
def purchases():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "auto":
            n = suggest_purchases()
            audit("Формирование плана закупок", "purchases", "", f"добавлено позиций: {n}")
            db().commit()
            flash(f"Добавлено позиций в план закупок: {n}", "info")
        elif action == "status":
            pid, st = to_int(request.form.get("id"), -1), request.form.get("status")
            p = q("SELECT * FROM purchases WHERE id=?", (pid,), one=True)
            if p and st in ("Запланирована", "Заказана", "Получена", "Отменена") and p["status"] != st:
                if p["status"] == "Получена":
                    flash("Закупка уже оформлена как полученная — статус не меняется", "warning")
                else:
                    x("UPDATE purchases SET status=? WHERE id=?", (st, pid))
                    if st == "Получена" and p["item_id"]:  # закрытие цикла: возврат в эксплуатацию
                        x("UPDATE items SET qty=qty+?, status='Исправно' WHERE id=?", (p["qty"], p["item_id"]))
                        log_move(p["item_id"], "Поступление", p["qty"], f"Закупка №{pid}: {p['reason']}")
                    audit("Изменён статус закупки", "purchases", pid, f"{p['status']} → {st}")
                    db().commit()
        return redirect(url_for("purchases"))
    return render_template("purchases.html", rows=q("SELECT * FROM purchases ORDER BY id DESC"),
                           states=["Запланирована", "Заказана", "Получена", "Отменена"])


# ---------------------------------------------------------------- справочники
DICTS = {
    "categories": dict(title="Категории", table="categories", cols=[("name", "Название", "text", True)],
                       text_refs=[("items", "category")], id_refs=[], protected=[]),
    "states": dict(title="Состояния", table="states", cols=[("name", "Название", "text", True)],
                   text_refs=[("items", "status"), ("checks", "status"), ("issues", "return_state")],
                   id_refs=[], protected=SYSTEM_STATES),
    "locations": dict(title="Места хранения", table="locations", cols=[("name", "Название", "text", True)],
                      text_refs=[("items", "location")], id_refs=[], protected=[]),
    "employees": dict(title="Сотрудники", table="employees",
                      cols=[("full_name", "ФИО", "text", True), ("position", "Должность", "text", False),
                            ("phone", "Телефон", "text", False), ("email", "E-mail", "email", False),
                            ("active", "Работает", "bool", False)],
                      text_refs=[], protected=[],
                      id_refs=[("issues", "employee_id"), ("bookings", "employee_id"),
                               ("users", "employee_id"), ("movements", "employee_id")]),
    "events": dict(title="Мероприятия", table="events",
                   cols=[("name", "Название", "text", True), ("event_date", "Дата", "date", False),
                         ("place", "Площадка", "text", False), ("note", "Примечание", "text", False)],
                   text_refs=[], protected=[NO_EVENT],
                   id_refs=[("issues", "event_id"), ("bookings", "event_id"), ("movements", "event_id")]),
}


def dict_or_404(name):
    if name not in DICTS:
        abort(404)
    return DICTS[name]


def dict_usage(d, row):
    key = d["cols"][0][0]
    n = sum(q(f"SELECT COUNT(*) c FROM {t} WHERE {c}=?", (row[key],), one=True)["c"] for t, c in d["text_refs"])
    return n + sum(q(f"SELECT COUNT(*) c FROM {t} WHERE {c}=?", (row["id"],), one=True)["c"] for t, c in d["id_refs"])


def dict_values(d, f):
    vals, err = {}, []
    for col, label, typ, req in d["cols"]:
        v = f.get(col, "").strip()
        if typ == "bool":
            v = 1 if f.get(col) else 0
        elif req and not v:
            err.append(f"Поле «{label}» обязательно")
        elif typ == "date" and v and not valid_date(v):
            err.append(f"Поле «{label}»: некорректная дата")
        elif typ == "email" and v and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", v):
            err.append(f"Поле «{label}»: некорректный e-mail")
        vals[col] = v
    return vals, err


@app.route("/dict/<name>", methods=["GET", "POST"])
@need("dictionaries")
def dictionary(name):
    d = dict_or_404(name)
    if request.method == "POST":
        vals, err = dict_values(d, request.form)
        if not err:
            cols = list(vals)
            try:
                cur = x(f"INSERT INTO {d['table']}({','.join(cols)}) VALUES({','.join('?' * len(cols))})",
                        [vals[c] for c in cols])
                audit("Справочник: добавлена запись", d["table"], cur.lastrowid, f"{d['title']}: {vals[cols[0]]}")
                db().commit()
                flash("Запись добавлена", "success")
                return redirect(url_for("dictionary", name=name))
            except sqlite3.IntegrityError:
                err.append("Такая запись уже существует")
        for e in err:
            flash(e, "danger")
    order = d["cols"][0][0]
    rows = q(f"SELECT * FROM {d['table']} ORDER BY " + ("COALESCE(event_date,'') DESC, name" if name == "events" else order))
    for r in rows:
        r["used"] = dict_usage(d, r)
        r["locked"] = r[order] in d["protected"]
    return render_template("dict.html", d=d, name=name, rows=rows, dicts=DICTS, order=order,
                           f=request.form if request.method == "POST" else {"active": "1"})


@app.route("/dict/<name>/<int:rid>/edit", methods=["GET", "POST"])
@need("dictionaries")
def dictionary_edit(name, rid):
    d = dict_or_404(name)
    row = q(f"SELECT * FROM {d['table']} WHERE id=?", (rid,), one=True)
    if not row:
        abort(404)
    key = d["cols"][0][0]
    locked = row[key] in d["protected"]
    if request.method == "POST":
        vals, err = dict_values(d, request.form)
        if locked:
            err.append("Системная запись: изменение и удаление запрещены, на неё опирается логика приложения")
        if not err:
            try:
                x(f"UPDATE {d['table']} SET " + ",".join(f"{c}=?" for c in vals) + " WHERE id=?", [*vals.values(), rid])
                if vals[key] != row[key]:  # переименование распространяем на карточки/акты
                    for t, c in d["text_refs"]:
                        x(f"UPDATE {t} SET {c}=? WHERE {c}=?", (vals[key], row[key]))
                audit("Справочник: изменена запись", d["table"], rid, f"{d['title']}: {row[key]} → {vals[key]}")
                db().commit()
                flash("Запись сохранена", "success")
                return redirect(url_for("dictionary", name=name))
            except sqlite3.IntegrityError:
                db().rollback()
                err.append("Такая запись уже существует")
        for e in err:
            flash(e, "danger")
    return render_template("dict_edit.html", d=d, name=name, row=row, locked=locked,
                           f=request.form if request.method == "POST" else row)


@app.route("/dict/<name>/<int:rid>/delete", methods=["POST"])
@need("dictionaries")
def dictionary_delete(name, rid):
    d = dict_or_404(name)
    row = q(f"SELECT * FROM {d['table']} WHERE id=?", (rid,), one=True)
    if row:
        key = d["cols"][0][0]
        if row[key] in d["protected"]:
            flash("Системную запись удалить нельзя", "danger")
        elif dict_usage(d, row):
            flash("Запись используется в учетных данных — удалить нельзя"
                  + (" (сотрудника можно отметить как неработающего)" if name == "employees" else ""), "danger")
        else:
            x(f"DELETE FROM {d['table']} WHERE id=?", (rid,))
            audit("Справочник: удалена запись", d["table"], rid, f"{d['title']}: {row[key]}")
            db().commit()
            flash("Запись удалена", "info")
    return redirect(url_for("dictionary", name=name))


# ---------------------------------------------------------------- инвентаризация
LINE_SQL = ("SELECT l.*, i.name, i.inv_no, i.category, i.location FROM inventory_lines l "
            "JOIN items i ON i.id=l.item_id WHERE l.inv_id=? ORDER BY i.location, i.name")


@app.route("/inventory", methods=["GET", "POST"])
@need("operations")
def inventory():
    if request.method == "POST":
        f = request.form
        loc, cat = f.get("location", ""), f.get("category", "")
        w, args = ["status<>'Списано'"], []
        for col, v in (("location", loc), ("category", cat)):
            if v:
                w.append(f"{col}=?")
                args.append(v)
        items_ = q(ITEM_SQL + " WHERE " + " AND ".join("i." + c for c in w), args)
        title = f.get("title", "").strip() or f"Инвентаризация от {today()}"
        scope = ", ".join(filter(None, [f"место: {loc}" if loc else "", f"категория: {cat}" if cat else ""])) or "все имущество"
        if not items_:
            flash("В выбранной области нет имущества", "danger")
        else:
            cur = x("INSERT INTO inventories(title,created,scope,created_by) VALUES(?,?,?,?)",
                    (title, now(), scope, g.user["id"]))
            for i in items_:
                x("INSERT INTO inventory_lines(inv_id,item_id,expected,issued) VALUES(?,?,?,?)",
                  (cur.lastrowid, i["id"], i["qty"] - i["issued"], i["issued"]))
            audit("Начата инвентаризация", "inventories", cur.lastrowid, f"{title} ({scope}), позиций: {len(items_)}")
            db().commit()
            return redirect(url_for("inventory_view", inv_id=cur.lastrowid))
    rows = q("SELECT v.*, (SELECT COUNT(*) FROM inventory_lines WHERE inv_id=v.id) n, "
             "(SELECT COUNT(*) FROM inventory_lines WHERE inv_id=v.id AND actual IS NOT NULL) done, "
             "(SELECT COUNT(*) FROM inventory_lines WHERE inv_id=v.id AND actual IS NOT NULL AND actual<>expected) diff "
             "FROM inventories v ORDER BY v.id DESC")
    return render_template("inventory.html", rows=rows, cats=names("categories"), locs=names("locations"))


@app.route("/inventory/<int:inv_id>", methods=["GET", "POST"])
@need("operations")
def inventory_view(inv_id):
    inv = q("SELECT * FROM inventories WHERE id=?", (inv_id,), one=True)
    if not inv:
        abort(404)
    lines = q(LINE_SQL, (inv_id,))
    if request.method == "POST" and inv["status"] == "Идёт":
        f, err = request.form, []
        for l in lines:
            raw = f.get(f"a{l['id']}", "").strip()
            l["new_actual"] = None if raw == "" else to_int(raw)
            if raw != "" and (l["new_actual"] is None or l["new_actual"] < 0):
                err.append(f"«{l['name']}»: фактическое количество — целое число ≥ 0")
            l["new_note"] = f.get(f"n{l['id']}", "").strip()
        action = f.get("action", "save")
        if not err and action != "save" and any(l["new_actual"] is None for l in lines):
            err.append(f"Для завершения заполните фактическое количество по всем позициям "
                       f"(не заполнено: {sum(l['new_actual'] is None for l in lines)})")
        if err:
            for e in err[:5]:
                flash(e, "danger")
        else:
            for l in lines:
                x("UPDATE inventory_lines SET actual=?, note=? WHERE id=?", (l["new_actual"], l["new_note"], l["id"]))
            if action == "save":
                audit("Сохранены результаты инвентаризации", "inventories", inv_id)
                db().commit()
                flash("Результаты сохранены", "success")
            else:
                apply = action == "finish_apply"
                fixed = 0
                for l in lines:
                    diff = l["new_actual"] - l["expected"]
                    if diff and apply:
                        x("UPDATE items SET qty=qty+? WHERE id=?", (diff, l["item_id"]))
                        log_move(l["item_id"], "Инвентаризация", diff,
                                 f"«{inv['title']}»: по учёту {l['expected']}, фактически {l['new_actual']} "
                                 f"→ корректировка {diff:+d}")
                        fixed += 1
                    elif diff:
                        log_move(l["item_id"], "Инвентаризация (расхождение)", diff,
                                 f"«{inv['title']}»: по учёту {l['expected']}, фактически {l['new_actual']}; без корректировки")
                    else:
                        log_move(l["item_id"], "Инвентаризация (сверка)", l["new_actual"], f"«{inv['title']}»: расхождений нет")
                x("UPDATE inventories SET status='Завершена', finished=?, applied=? WHERE id=?",
                  (now(), 1 if apply else 0, inv_id))
                audit("Завершена инвентаризация", "inventories", inv_id,
                      f"{inv['title']}; корректировок: {fixed}" if apply else f"{inv['title']}; без корректировки")
                db().commit()
                flash("Инвентаризация завершена" + (f", скорректировано позиций: {fixed}" if apply else ""), "success")
            return redirect(url_for("inventory_view", inv_id=inv_id))
        db().rollback()
    inv = q("SELECT * FROM inventories WHERE id=?", (inv_id,), one=True)
    lines = q(LINE_SQL, (inv_id,))
    return render_template("inventory_view.html", inv=inv, lines=lines)


# ---------------------------------------------------------------- отчеты
def fmt_value(v, kind):
    if v is None or v == "":
        return ""
    if kind == "int":
        return f"{int(v):,}".replace(",", "\u00a0")
    if kind == "money":
        return f"{float(v):,.2f}".replace(",", "\u00a0").replace(".", ",")
    return str(v)


def finish_rep(rep, a):
    rep["disp"] = [[fmt_value(v, k) for v, k in zip(r, rep["kinds"])] for r in rep["rows"]]
    rep["totals_disp"] = [fmt_value(v, k) for v, k in zip(rep["totals"], rep["kinds"])] if rep.get("totals") else None
    rep.setdefault("vals", dict(a))
    parts = []
    for name, label, typ, opts in rep["filters"]:
        v = rep["vals"].get(name)
        if v:
            if typ == "emp":
                v = next((o["full_name"] for o in opts if str(o["id"]) == v), v)
            parts.append(f"{label}: {'да' if typ == 'check' else v}")
    rep["subtitle"] = "; ".join(parts) if parts else "без ограничений"
    return rep


def rep_stock(a):
    rows = search_items(a)
    data = [[r["inv_no"], r["name"], r["category"], r["location"], r["status"], r["avail"], r["qty"], r["issued"],
             r["free"], r["price"], r["price"] * r["qty"]] for r in rows]
    tot = ["", "Итого", "", "", "", "", sum(r[6] for r in data), sum(r[7] for r in data), sum(r[8] for r in data),
           None, sum(r[10] for r in data)]
    c = item_filters_ctx()
    return finish_rep(dict(
        title="Наличие имущества", headers=["Инв. №", "Наименование", "Категория", "Место хранения", "Состояние",
                                            "Статус", "Всего", "Выдано", "В наличии", "Цена, ₽", "Стоимость, ₽"],
        kinds=["text"] * 6 + ["int"] * 3 + ["money"] * 2, widths=[8, 22, 14, 13, 12, 12, 7, 8, 9, 13, 15],
        rows=data, totals=tot,
        filters=[("cat", "Категория", "select", c["cats"]), ("loc", "Место хранения", "select", c["locs"]),
                 ("status", "Статус", "select", c["avail"]), ("state", "Состояние", "select", c["states"]),
                 ("s", "Поиск", "text", None)]), a)


def rep_movements(a):
    df = a["date_from"] if "date_from" in a else (date.today() - timedelta(days=30)).isoformat()
    dt = a.get("date_to", "")
    w, args = [], []
    if valid_date(df):
        w.append("substr(m.ts,1,10)>=?")
        args.append(df)
    if valid_date(dt):
        w.append("substr(m.ts,1,10)<=?")
        args.append(dt)
    for key, col in (("kind", "m.kind"), ("cat", "i.category")):
        if a.get(key):
            w.append(f"{col}=?")
            args.append(a[key])
    rows = q(MOVE_SQL + (" WHERE " + " AND ".join(w) if w else "") + " ORDER BY m.id DESC LIMIT 5000", args)
    s = (a.get("s") or "").strip().casefold()
    if s:
        rows = [r for r in rows if s in f"{r['item_name']} {r['inv_no']}".casefold()]
    data = [[r["ts"][:16], r["inv_no"], r["item_name"], r["kind"], r["qty"], r["emp_name"] or "", r["event_name"] or "",
             r["user_name"] or "", r["details"] or ""] for r in rows]
    kinds = [r["kind"] for r in q("SELECT DISTINCT kind FROM movements ORDER BY kind")]
    return finish_rep(dict(
        title="Движение имущества", headers=["Дата и время", "Инв. №", "Наименование", "Операция", "Кол-во",
                                             "Сотрудник", "Мероприятие", "Выполнил", "Примечание"],
        kinds=["text"] * 4 + ["int"] + ["text"] * 4, widths=[15, 8, 24, 16, 6, 18, 20, 16, 34], rows=data,
        totals=["", "", f"Операций: {len(data)}", "", None, "", "", "", ""],
        vals={**a, "date_from": df},
        filters=[("date_from", "С даты", "date", None), ("date_to", "По дату", "date", None),
                 ("kind", "Операция", "select", kinds), ("cat", "Категория", "select", names("categories")),
                 ("s", "Поиск", "text", None)]), a)


def rep_debts(a):
    w, args = ["s.returned_at IS NULL"], []
    if a.get("emp"):
        w.append("s.employee_id=?")
        args.append(to_int(a["emp"], -1))
    if a.get("overdue"):
        w.append("s.due_date<?")
        args.append(today())
    rows = [decorate_issue(r) for r in q(ISSUE_SQL + " WHERE " + " AND ".join(w) + " ORDER BY e.full_name, s.due_date", args)]
    data = [[r["emp_name"], r["inv_no"], r["item_name"], r["qty"], r["event_name"] or "", r["issued_at"],
             r["due_date"], max(0, -r["days"]), r["issued_by_name"] or ""] for r in rows]
    return finish_rep(dict(
        title="Задолженности по возврату имущества",
        headers=["Сотрудник", "Инв. №", "Наименование", "Кол-во", "Мероприятие", "Выдано", "Вернуть до",
                 "Просрочка, дн.", "Выдал"],
        kinds=["text"] * 3 + ["int"] + ["text"] * 3 + ["int", "text"], widths=[22, 8, 25, 7, 24, 11, 11, 11, 17],
        rows=data, totals=["Итого", "", f"Записей: {len(data)}", sum(r[3] for r in data), "", "", "", None, ""],
        filters=[("emp", "Сотрудник", "emp", q("SELECT id, full_name FROM employees ORDER BY full_name")),
                 ("overdue", "Только просроченные", "check", None)]), a)


def rep_categories(a):
    agg = {n: dict(pos=0, qty=0, out=0, ok=0, rep=0, off=0, val=0.0) for n in names("categories")}
    for i in enrich(q(ITEM_SQL)):
        c = agg.setdefault(i["category"] or "(без категории)", dict(pos=0, qty=0, out=0, ok=0, rep=0, off=0, val=0.0))
        c["pos"] += 1
        if i["status"] == "Списано":
            c["off"] += 1
            continue
        c["qty"] += i["qty"]
        c["out"] += i["issued"]
        c["val"] += i["price"] * i["qty"]
        c["rep" if i["stage"] == "Ремонт" else "ok"] += 1
    data = [[n, c["pos"], c["qty"], c["out"], c["qty"] - c["out"], c["ok"], c["rep"], c["off"], c["val"]]
            for n, c in sorted(agg.items())]
    tot = ["Итого"] + [sum(r[i] for r in data) for i in range(1, 9)]
    return finish_rep(dict(
        title="Сводка по категориям", headers=["Категория", "Позиций", "Единиц (без списанных)", "Выдано", "В наличии",
                                               "Исправно, поз.", "Ремонт, поз.", "Списано, поз.", "Стоимость, ₽"],
        kinds=["text"] + ["int"] * 7 + ["money"], widths=[26, 9, 14, 9, 10, 10, 10, 10, 17], rows=data, totals=tot,
        filters=[]), a)


def rep_inventory(inv_id):
    inv = q("SELECT * FROM inventories WHERE id=?", (inv_id,), one=True)
    if not inv:
        abort(404)
    lines = q(LINE_SQL, (inv_id,))
    data = [[l["inv_no"], l["name"], l["location"], l["expected"], l["issued"], l["actual"],
             (l["actual"] - l["expected"]) if l["actual"] is not None else None, l["note"] or ""] for l in lines]
    diff = sum(1 for r in data if r[6])
    rep = dict(
        title=f"Сличительная ведомость: {inv['title']}",
        headers=["Инв. №", "Наименование", "Место хранения", "По учёту (на месте)", "Выдано (у сотрудников)",
                 "Фактически", "Расхождение", "Примечание"],
        kinds=["text"] * 3 + ["int"] * 4 + ["text"], widths=[9, 30, 16, 12, 12, 11, 11, 30], rows=data,
        totals=["", f"Позиций: {len(data)}, с расхождениями: {diff}", "", None, None, None, None, ""], filters=[])
    finish_rep(rep, {})
    rep["subtitle"] = f"Область: {inv['scope']}; начата {inv['created'][:16]}; статус: {inv['status']}"
    return rep


REPORTS = {
    "stock": dict(title="Наличие имущества", fn=rep_stock, hint="Остатки, выдано и свободно, стоимость. Фильтры: категория, место, статус, состояние."),
    "movements": dict(title="Движение имущества", fn=rep_movements, hint="Постановка на учёт, выдачи, возвраты, перемещения, корректировки за период."),
    "debts": dict(title="Задолженности", fn=rep_debts, hint="Невозвращённое имущество по сотрудникам, просрочка в днях."),
    "categories": dict(title="По категориям", fn=rep_categories, hint="Сводка: количество, выдано, состояние, стоимость фонда."),
}


_FONT = None


def pdf_fonts():
    global _FONT
    if _FONT:
        return _FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    from reportlab.pdfbase.ttfonts import TTFont
    fd = os.path.join(BASE_DIR, "fonts")
    cands = [(os.path.join(fd, "DejaVuSans.ttf"), os.path.join(fd, "DejaVuSans-Bold.ttf")),
             ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
             ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
             (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
             ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf")]
    for reg, bold in cands:
        if os.path.exists(reg):
            pdfmetrics.registerFont(TTFont("AppFont", reg))
            pdfmetrics.registerFont(TTFont("AppFont-Bold", bold if os.path.exists(bold) else reg))
            registerFontFamily("AppFont", normal="AppFont", bold="AppFont-Bold", italic="AppFont", boldItalic="AppFont-Bold")
            _FONT = ("AppFont", "AppFont-Bold")
            return _FONT
    raise RuntimeError("Не найден шрифт с кириллицей для PDF. Поместите DejaVuSans.ttf и DejaVuSans-Bold.ttf в папку fonts/")


def export_xlsx(rep):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", " ", rep["title"])[:31]
    n = len(rep["headers"])
    ws.cell(1, 1, rep["title"]).font = Font(bold=True, size=14)
    ws.cell(2, 1, f"Фильтры: {rep['subtitle']}").font = Font(italic=True, color="555555")
    ws.cell(3, 1, f"Сформировано: {now()[:16]}" + (f", {g.user['full_name']}" if g.get("user") else "")).font = Font(color="555555")
    hr = 5
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(rep["headers"], 1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B1A2B")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border

    def put(r, vals, bold=False):
        for c, (v, k) in enumerate(zip(vals, rep["kinds"]), 1):
            cell = ws.cell(r, c, v)
            if isinstance(v, str) and v.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"  # текст пользователя не должен выполняться как формула
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right" if k in ("int", "money") else "left")
            if k == "int":
                cell.number_format = "#,##0"
            elif k == "money":
                cell.number_format = "#,##0.00"
            if bold:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="F1E4E8")
    for i, row in enumerate(rep["rows"]):
        put(hr + 1 + i, row)
    last = hr + len(rep["rows"])
    if rep.get("totals"):
        put(last + 1, rep["totals"], True)
    for c, wd in enumerate(rep["widths"], 1):
        ws.column_dimensions[get_column_letter(c)].width = wd + 2
    ws.freeze_panes = ws.cell(hr + 1, 1)
    if rep["rows"]:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(n)}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_pdf(rep):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    reg, bold = pdf_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=14 * mm, title=rep["title"], author="Система учета АГФ")
    st = ParagraphStyle("c", fontName=reg, fontSize=7.5, leading=9.5)
    stn = ParagraphStyle("n", parent=st, alignment=2)
    sth = ParagraphStyle("h", parent=st, fontName=bold, textColor=colors.white, alignment=1)
    stb = ParagraphStyle("b", parent=st, fontName=bold)
    stbn = ParagraphStyle("bn", parent=stb, alignment=2)

    def cell(v, kind, head=False, tot=False):
        s = xml_escape(str(v)).replace("\n", "<br/>")
        if head:
            return Paragraph(s, sth)
        num = kind in ("int", "money")
        return Paragraph(s, (stbn if num else stb) if tot else (stn if num else st))

    width = doc.width
    wsum = sum(rep["widths"])
    colw = [width * w / wsum for w in rep["widths"]]
    data = [[cell(h, "text", head=True) for h in rep["headers"]]]
    data += [[cell(v, k) for v, k in zip(r, rep["kinds"])] for r in rep["disp"]]
    if rep.get("totals_disp"):
        data.append([cell(v, k, tot=True) for v, k in zip(rep["totals_disp"], rep["kinds"])])
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5B1A2B")),
             ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F2F3")]),
             ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
    if rep.get("totals_disp"):
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F1E4E8")))
    tbl = Table(data, colWidths=colw, repeatRows=1)
    tbl.setStyle(TableStyle(style))
    who = g.user["full_name"] if g.get("user") else ""
    head = [Paragraph(xml_escape(rep["title"]), ParagraphStyle("t", fontName=bold, fontSize=14, leading=18)),
            Paragraph(xml_escape(f"Фильтры: {rep['subtitle']}"), ParagraphStyle("s", parent=st, fontSize=8.5, textColor=colors.HexColor("#555555"))),
            Spacer(1, 4 * mm)]

    def footer(canv, d):
        canv.saveState()
        canv.setFont(reg, 7)
        canv.setFillColor(colors.HexColor("#666666"))
        canv.drawString(10 * mm, 7 * mm, f"Астраханская государственная филармония · сформировано {now()[:16]} {who}")
        canv.drawRightString(landscape(A4)[0] - 10 * mm, 7 * mm, f"Страница {d.page}")
        canv.restoreState()
    doc.build(head + [tbl], onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf


def export_response(rep, fmt, key):
    try:
        buf = export_xlsx(rep) if fmt == "xlsx" else export_pdf(rep)
    except RuntimeError as e:
        flash(str(e), "danger")
        return None
    audit("Экспорт отчёта", "reports", key, f"{rep['title']} → {fmt.upper()}; фильтры: {rep['subtitle']}")
    db().commit()
    mime = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if fmt == "xlsx" else "application/pdf")
    return send_file(buf, mimetype=mime, as_attachment=True, download_name=f"report_{key}_{today()}.{fmt}")


@app.route("/reports")
@need("reports")
def reports():
    return render_template("reports.html", reports=REPORTS)


@app.route("/reports/<key>")
@need("reports")
def report(key):
    if key not in REPORTS:
        abort(404)
    rep = REPORTS[key]["fn"](request.args)
    fmt = request.args.get("fmt")
    if fmt in ("xlsx", "pdf"):
        resp = export_response(rep, fmt, key)
        if resp:
            return resp
    qs = {k: v for k, v in request.args.items() if k != "fmt"}
    return render_template("report.html", rep=rep, key=key, qs=qs, reports=REPORTS)


@app.route("/inventory/<int:inv_id>/export")
@need("reports")
def inventory_export(inv_id):
    fmt = request.args.get("fmt", "xlsx")
    if fmt not in ("xlsx", "pdf"):
        abort(400)
    resp = export_response(rep_inventory(inv_id), fmt, f"inventory_{inv_id}")
    return resp or redirect(url_for("inventory_view", inv_id=inv_id))


# ---------------------------------------------------------------- администрирование: пользователи
LOGIN_RE = re.compile(r"[A-Za-z0-9_.\-]{3,32}")


def other_admins(uid):
    return q("SELECT COUNT(*) c FROM users WHERE role='admin' AND active=1 AND id<>?", (uid,), one=True)["c"]


@app.route("/admin/users", methods=["GET", "POST"])
@need("admin")
def admin_users():
    if request.method == "POST":
        f, err = request.form, []
        lg, pw, role = f.get("login", "").strip(), f.get("password", ""), f.get("role", "")
        emp = to_int(f.get("employee_id")) if f.get("employee_id") else None
        if not LOGIN_RE.fullmatch(lg):
            err.append("Логин: 3–32 символа, латиница, цифры и . _ -")
        if not f.get("full_name", "").strip():
            err.append("Укажите ФИО")
        if role not in ROLES:
            err.append("Выберите роль")
        pe = validate_password(pw)
        if pe:
            err.append(pe)
        if emp and not q("SELECT 1 FROM employees WHERE id=?", (emp,), one=True):
            err.append("Сотрудник не найден")
        if not err:
            try:
                cur = x("INSERT INTO users(login,pw_hash,full_name,role,employee_id,created) VALUES(?,?,?,?,?,?)",
                        (lg, generate_password_hash(pw), f["full_name"].strip(), role, emp, now()))
                audit("Создан пользователь", "users", cur.lastrowid, f"{lg}, роль: {ROLES[role]}")
                db().commit()
                flash("Пользователь создан", "success")
                return redirect(url_for("admin_users"))
            except sqlite3.IntegrityError:
                err.append("Пользователь с таким логином уже существует")
        for e in err:
            flash(e, "danger")
    rows = q("SELECT u.*, e.full_name AS emp_name FROM users u LEFT JOIN employees e ON e.id=u.employee_id ORDER BY u.id")
    return render_template("admin_users.html", rows=rows, emps=q("SELECT id, full_name FROM employees ORDER BY full_name"),
                           f=request.form if request.method == "POST" else {})


@app.route("/admin/users/<int:uid>", methods=["GET", "POST"])
@need("admin")
def admin_user_edit(uid):
    u = q("SELECT * FROM users WHERE id=?", (uid,), one=True)
    if not u:
        abort(404)
    if request.method == "POST":
        f, err = request.form, []
        role, active = f.get("role", ""), 1 if f.get("active") else 0
        emp = to_int(f.get("employee_id")) if f.get("employee_id") else None
        if role not in ROLES:
            err.append("Выберите роль")
        if not f.get("full_name", "").strip():
            err.append("Укажите ФИО")
        if uid == g.user["id"] and not active:
            err.append("Нельзя отключить собственную учётную запись")
        if (role != "admin" or not active) and u["role"] == "admin" and u["active"] and not other_admins(uid):
            err.append("Должен остаться хотя бы один активный администратор")
        if f.get("password") and validate_password(f["password"]):
            err.append(validate_password(f["password"]))
        if not err:
            x("UPDATE users SET full_name=?, role=?, employee_id=?, active=? WHERE id=?",
              (f["full_name"].strip(), role, emp, active, uid))
            changes = []
            if role != u["role"]:
                changes.append(f"роль: {ROLES[u['role']]} → {ROLES[role]}")
            if active != u["active"]:
                changes.append("доступ " + ("включён" if active else "отключён"))
            if f.get("password"):
                x("UPDATE users SET pw_hash=? WHERE id=?", (generate_password_hash(f["password"]), uid))
                changes.append("сброшен пароль")
            audit("Изменён пользователь", "users", uid, f"{u['login']}: " + ("; ".join(changes) or "данные профиля"))
            db().commit()
            flash("Сохранено", "success")
            return redirect(url_for("admin_users"))
        for e in err:
            flash(e, "danger")
    return render_template("admin_user_edit.html", u=u, f=request.form if request.method == "POST" else u,
                           emps=q("SELECT id, full_name FROM employees ORDER BY full_name"))


@app.route("/issues/<int:sid>/photos")
@need("view")
def issue_return_photos(sid):
    row = q("SELECT return_photos, item_id FROM issues WHERE id=?", (sid,), one=True)
    if not row or not row["return_photos"]:
        abort(404)
    try:
        photos = json.loads(row["return_photos"] or "[]")
    except (TypeError, ValueError):
        photos = []
    return render_template("return_photos.html", photos=photos, sid=sid, item=get_item(row["item_id"]))


@app.route("/return-photos/<path:filename>")
@need("view")
def return_photo(filename):
    safe = os.path.basename(filename)
    if safe != filename or not re.fullmatch(r"[a-f0-9]{32}\.(?:jpg|jpeg|png|webp)", safe):
        abort(404)
    return send_from_directory(UPLOAD_DIR, safe)


# Шаблоны вынесены в отдельный модуль для разделения интерфейса и логики.
from templates_ui import T
app.jinja_loader = DictLoader(T)

# ---------------------------------------------------------------- запуск
init_db()


if __name__ == "__main__":
    import sys
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "remind":      # python app.py remind  — для планировщика (cron / Планировщик заданий)
        with app.app_context():
            sent, skipped, err = send_reminders()
            con = sqlite3.connect(DB, timeout=15)
            write_audit(con, None, "system", "Отправка напоминаний (по расписанию)", "notifications", "",
                        f"отправлено: {sent}, без e-mail: {skipped}" + (f"; {err}" if err else ""))
            con.commit()
            con.close()
        print(f"Отправлено: {sent}, без e-mail: {skipped}" + (f"; {err}" if err else ""))
    else:
        debug = os.environ.get("APP_DEBUG", "1") == "1"
        app.run(debug=debug)
